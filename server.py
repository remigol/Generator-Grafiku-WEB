import os,json,calendar,datetime,tempfile
from http.server import ThreadingHTTPServer,SimpleHTTPRequestHandler
from solver import solve,validate,build_days,PEOPLE
from openpyxl import Workbook
BASE=os.path.dirname(os.path.abspath(__file__))
class Handler(SimpleHTTPRequestHandler):
 def __init__(self,*a,**kw):super().__init__(*a,directory=BASE,**kw)
 def do_POST(self):
  try:
   n=int(self.headers.get('Content-Length','0'))
   if n>200000:raise ValueError('Za duży formularz')
   data=json.loads(self.rfile.read(n));path=self.path.split('?')[0]
   if path=='/api/generate':return self.json(solve(data,relax={'one_day_rest'} if data.get('oneDayRestTest') else None))
   if path=='/api/diagnose':return self.json(solve(data,diagnostic=True))
   if path=='/api/validate':return self.json({'issues':validate(data,allow_shortfall=bool(data.get('partial')),allow_one_day_rest=bool(data.get('oneDayRestTest'))),'shortfalls':[{'employee':p,'hours':12*sum(x in ('6','18') for x in data['schedule'][p]),'missing':max(0,12*__import__('math').ceil(int(data.get('nominal',176))/12)-12*sum(x in ('6','18') for x in data['schedule'][p]))} for p in PEOPLE]})
   if path=='/api/export':
    issues=validate(data,allow_shortfall=bool(data.get('partial')),allow_one_day_rest=bool(data.get('oneDayRestTest')))
    return self.export(data,issues)
   return self.json({'error':'Nieznana ścieżka'},404)
  except Exception as e:return self.json({'error':str(e)},400)
 def json(self,obj,status=200):
  b=json.dumps(obj,ensure_ascii=False).encode();self.send_response(status);self.send_header('Content-Type','application/json; charset=utf-8');self.send_header('Content-Length',str(len(b)));self.end_headers();self.wfile.write(b)
 def export(self,data,issues=None):
  days=build_days(int(data['year']),int(data['month']),set(map(int,data.get('closed',[]))))
  wb=Workbook();ws=wb.active;ws.title='Grafik miesięczny';D=len(days)
  from openpyxl.styles import PatternFill,Font,Alignment
  from openpyxl.utils import get_column_letter
  ws.append(['Data','Dzień tygodnia']+PEOPLE+['Obsada 6','Obsada 18'])
  weekdays=('Pn','Wt','Śr','Cz','Pt','So','Nd')
  for i,day in enumerate(days):
   date=datetime.date(int(data['year']),int(data['month']),day['day'])
   row=[date,weekdays[day['weekday']]]+[data['schedule'][p][i] or '' for p in PEOPLE]
   row.extend([sum(data['schedule'][p][i]=='6' for p in PEOPLE),sum(data['schedule'][p][i]=='18' for p in PEOPLE)])
   ws.append(row)
   ws.cell(i+2,1).number_format='DD.MM.YYYY'
   if day['type']!='weekday':
    bg='F2E8FA' if day['type']=='closed' else 'FFF0C9'
    for cell in ws[i+2]:cell.fill=PatternFill('solid',fgColor=bg)
   for j,p in enumerate(PEOPLE,3):
    cell=ws.cell(i+2,j)
    if cell.value=='6':cell.fill=PatternFill('solid',fgColor='DDF2E5')
    elif cell.value=='18':cell.fill=PatternFill('solid',fgColor='E4E7FA')
  for cell in ws[1]:
   cell.fill=PatternFill('solid',fgColor='17324E')
   cell.font=Font(color='FFFFFF',bold=True)
  ws.freeze_panes='C2'
  ws.auto_filter.ref=f'A1:M{D+1}'
  ws.column_dimensions['A'].width=15
  ws.column_dimensions['B'].width=17
  for j in range(3,12):ws.column_dimensions[get_column_letter(j)].width=9
  ws.column_dimensions['L'].width=13
  ws.column_dimensions['M'].width=14
  for row in ws.iter_rows(min_row=2):
   for cell in row[1:]:cell.alignment=Alignment(horizontal='center')
  summary=wb.create_sheet('Podsumowanie')
  summary.append(['Pracownik','Godziny','Zm. 6','Zm. 18','Nominał','Nadgodziny'])
  for p in PEOPLE:
   v=data['schedule'][p];h=12*sum(x in ('6','18') for x in v);n=int(data.get('nominal',176))
   summary.append([p,h,v.count('6'),v.count('18'),n,h-n])
  for cell in summary[1]:
   cell.fill=PatternFill('solid',fgColor='17324E')
   cell.font=Font(color='FFFFFF',bold=True)
  summary.freeze_panes='B2'
  for col in 'ABCDEF':summary.column_dimensions[col].width=17
  info=wb.create_sheet('Informacje')
  for row in [['Generator Grafiku WEB','v1'],['Miesiąc',str(data['month'])+'/'+str(data['year'])],['Nominał',int(data.get('nominal',176))],['Uwaga','Grafik z niedoborami godzin' if data.get('partial') else 'Format importu ShiftMate wymaga potwierdzenia']]:info.append(row)
  if issues:
   warnings=wb.create_sheet('Ostrzeżenia')
   warnings.append(['UWAGA: Grafik roboczy – wykryto naruszenia zasad. Nie używać bez weryfikacji.'])
   for issue in issues:warnings.append([issue])
   warnings.column_dimensions['A'].width=110
  import io
  out=io.BytesIO();wb.save(out);b=out.getvalue()
  self.send_response(200);self.send_header('Content-Type','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet');self.send_header('Content-Disposition','attachment; filename="Grafik_%s_%02d.xlsx"'%(data['year'],int(data['month'])));self.send_header('Content-Length',str(len(b)));self.end_headers();self.wfile.write(b)
if __name__=='__main__':
 port=int(os.environ.get('PORT','8000'));ThreadingHTTPServer(('0.0.0.0',port),Handler).serve_forever()
