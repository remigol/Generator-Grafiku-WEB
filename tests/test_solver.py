import ast
import calendar
import unittest
from pathlib import Path
from solver import solve, validate, build_days, PEOPLE

class SolverTests(unittest.TestCase):
    def test_export_manual_edits_with_warnings(self):
        import io
        from openpyxl import load_workbook
        from server import Handler
        class Dummy:
            def __init__(self):self.wfile=io.BytesIO();self.status=None
            def send_response(self,status):self.status=status
            def send_header(self,*args):pass
            def end_headers(self):pass
        data={"year":2026,"month":11,"nominal":160,"closed":[],"schedule":{p:[""]*30 for p in PEOPLE}}
        data["schedule"]["A"][1]="6"
        issues=["A: testowe naruszenie zasad"]
        dummy=Dummy()
        Handler.export(dummy,data,issues)
        self.assertEqual(dummy.status,200)
        wb=load_workbook(io.BytesIO(dummy.wfile.getvalue()))
        self.assertEqual(wb["Grafik miesięczny"]["A1"].value,"Data")
        self.assertEqual(wb["Grafik miesięczny"]["B1"].value,"Dzień tygodnia")
        self.assertEqual(wb["Grafik miesięczny"]["C1"].value,"A")
        self.assertEqual(wb["Grafik miesięczny"]["K1"].value,"I")
        self.assertEqual(wb["Grafik miesięczny"]["B3"].value,"Pn")
        self.assertEqual(wb["Grafik miesięczny"]["C3"].value,"6")
        self.assertEqual(wb["Grafik miesięczny"]["L3"].value,1)
        self.assertEqual(wb["Grafik miesięczny"]["M3"].value,0)
        self.assertEqual(wb["Grafik miesięczny"].max_row,31)
        self.assertIn("Podsumowanie",wb.sheetnames)
        self.assertEqual(wb["Ostrzeżenia"]["A2"].value,issues[0])
        self.assertIn("Informacje",wb.sheetnames)
    def test_syntax(self):
        ast.parse(Path("solver.py").read_text(encoding="utf-8"))
    def test_weekly_full_leave_validation(self):
        data={"year":2026,"month":11,"nominal":160,"closed":[],"locks":{},"weeklyShifts":{"2026-11-02":{"A":"W"}}}
        schedule={p:[""]*30 for p in PEOPLE}
        schedule["A"][1]="6"
        issues=validate({**data,"schedule":schedule},allow_shortfall=True)
        self.assertTrue(any("A, dzień 2:" in x and "wyborem W" in x for x in issues))
        schedule["A"][1]=""
        issues=validate({**data,"schedule":schedule},allow_shortfall=True)
        self.assertFalse(any("wyborem W" in x for x in issues))
    def test_weekly_full_leave_solver(self):
        data={"year":2026,"month":11,"nominal":160,"closed":[],"locks":{},"weeklyShifts":{"2026-11-02":{"A":"W"}}}
        result=solve(data,relax={"one_day_rest"})
        print("WEEK_LEAVE",result.get("ok"),result.get("message"),flush=True)
        if result.get("ok"):
            self.assertTrue(all(result["schedule"]["A"][d]=="" for d in range(1,8)))
            self.assertFalse(any("wyborem W" in x for x in validate({**data,"schedule":result["schedule"]},allow_one_day_rest=True)))
    def test_calendar_november(self):
        days=build_days(2026,11,set())
        self.assertEqual(len(days),30)
        self.assertEqual(sum(x["type"]=="sat" for x in days),4)
        self.assertEqual(sum(x["type"]=="weekday" for x in days),20)
        self.assertEqual(sum(x["type"]=="closed" for x in days),6)
    def test_november_2026_160(self):
        data={"year":2026,"month":11,"nominal":160,"closed":[],"locks":{},"weeklyShifts":{}}
        result=solve(data)
        print("NOVEMBER_RESULT",result.get("ok"),result.get("status"),result.get("message"),flush=True)
        self.assertTrue(result.get("ok"),result.get("message"))
        self.assertTrue(result.get("partial"))
        sched=result["schedule"]
        self.assertEqual(set(sched),set(PEOPLE))
        self.assertEqual(validate({**data,"schedule":sched},allow_shortfall=True),[])
        for p in PEOPLE:
            v=sched[p]
            self.assertEqual(len(v),30)
            self.assertLessEqual(sum(x in ("6","18") for x in v),14)
            self.assertLessEqual(abs(v.count("6")-v.count("18")),4)
            for i in range(1,29):
                self.assertFalse(v[i] in ("6","18") and v[i-1] not in ("6","18") and v[i+1] not in ("6","18"))
        self.assertGreater(sum(x["missing"] for x in result["shortfalls"]),0)
        self.assertGreaterEqual(sum(x["hours"] for x in result["shortfalls"]),117*12)
        self.assertLessEqual(max(x["missing"] for x in result["shortfalls"]),24)
        print("FAIRNESS_SHORTFALLS",{x["employee"]:x["missing"] for x in result["shortfalls"]},flush=True)
        print("NOVEMBER_DISTRIBUTION",{p:(sched[p].count("6"),sched[p].count("18")) for p in PEOPLE},flush=True)
    def test_max_equal_shifts_november(self):
        for shifts in range(14,7,-1):
            data={"year":2026,"month":11,"nominal":shifts*12,"closed":[],"locks":{},"weeklyShifts":{}}
            result=solve(data)
            print("CAPACITY_TEST",shifts,"shifts",shifts*12,"hours",result.get("ok"),result.get("status"),result.get("message"),flush=True)
            if result.get("ok") and not result.get("partial"):
                issues=validate({**data,"schedule":result["schedule"]})
                print("VALIDATION",len(issues),issues[:8],flush=True)
                self.assertEqual(issues,[])
                return
        self.fail("Nie znaleziono pełnego grafiku nawet dla 8 zmian na osobę")
    def test_single_day_blocks_november_168(self):
        data={"year":2026,"month":11,"nominal":168,"closed":[],"locks":{},"weeklyShifts":{}}
        result=solve(data,relax={"singletons"})
        print("SINGLETONS_168",result.get("ok"),result.get("status"),result.get("message"),flush=True)
        if result.get("ok"):
            schedule=result["schedule"]
            issues=validate({**data,"schedule":schedule})
            singleton_issues=[i for i in issues if "pojedynczy dzień pracy" in i]
            other_issues=[i for i in issues if i not in singleton_issues]
            self.assertEqual(other_issues,[])
            self.assertTrue(all(sum(v in ("6","18") for v in schedule[p])==14 for p in PEOPLE))
            print("SINGLETONS_COUNT",len(singleton_issues),flush=True)
        else:
            self.assertIn(result.get("status"),(1,2))
    def test_one_day_rest_november_168_diagnostic(self):
        from datetime import datetime, timedelta
        data={"year":2026,"month":11,"nominal":168,"closed":[],"locks":{},"weeklyShifts":{}}
        result=solve(data,relax={"one_day_rest"})
        print("ONE_DAY_REST_168",result.get("ok"),result.get("status"),result.get("message"),flush=True)
        if not result.get("ok"):
            self.assertIn(result.get("status"),(1,2))
            return
        schedule=result["schedule"]
        for d in range(30): print("SCHEDULE_ROW",d+1,",".join(schedule[p][d] or "-" for p in PEOPLE),flush=True)
        issues=validate({**data,"schedule":schedule})
        rest_issues=[i for i in issues if "tylko jeden dzień wolny" in i or "poniedziałek po sobocie" in i]
        self.assertEqual([i for i in issues if i not in rest_issues],[])
        self.assertTrue(all(sum(v in ("6","18") for v in schedule[p])==14 for p in PEOPLE))
        short_daily=[];weekly_gaps=[]
        for p in PEOPLE:
            shifts=[]
            for d,v in enumerate(schedule[p]):
                if v in ("6","18"):
                    start=datetime(2026,11,d+1,int(v))
                    shifts.append((start,start+timedelta(hours=12)))
            for prev,nxt in zip(shifts,shifts[1:]):
                gap=(nxt[0]-prev[1]).total_seconds()/3600
                if gap<11:short_daily.append((p,prev[0].day,nxt[0].day,gap))
            # Diagnostycznie: każda kolejna siedmiodniowa ramka musi
            # zawierać nieprzerwany okres odpoczynku >=35 h.
            for day in range(1,25):
                a=datetime(2026,11,day);b=a+timedelta(days=7)
                intervals=[(a,a)]+[(max(a,x),min(b,y)) for x,y in shifts if x<b and y>a]+[(b,b)]
                intervals.sort()
                max_gap=max((intervals[i+1][0]-intervals[i][1]).total_seconds()/3600 for i in range(len(intervals)-1))
                if max_gap<35:weekly_gaps.append((p,day,round(max_gap,1)))
        print("ONE_DAY_REST_VIOLATIONS",len(rest_issues),rest_issues[:12],flush=True)
        print("DAILY_REST_UNDER_11H",short_daily,flush=True)
        print("ROLLING_7DAY_UNDER_35H",len(weekly_gaps),weekly_gaps[:15],flush=True)
        self.assertEqual(short_daily,[])
    def test_manual_constraints_november_comparison(self):
        from datetime import datetime,timedelta
        base={"year":2026,"month":11,"nominal":168,"closed":[],"locks":{},"weeklyShifts":{}}
        # Dwa wolne dni na osobę: wybrane dni otwarcia, nie święta/niedziele.
        off={"A":[4,12],"B":[5,18],"C":[6,23],"D":[9,25],"E":[10,24],
             "F":[12,26],"G":[13,27],"H":[16,25],"I":[17,23]}
        locks={f"{p}-{d}":"W" for p,dates in off.items() for d in dates}
        weekly={
          "2026-11-02":{"A":"6","B":"6","C":"18"},
          "2026-11-09":{"D":"6","E":"6","G":"18"},
          "2026-11-16":{"F":"18","H":"18","I":"18"},
        }
        variants=[("BASE",base),("MANUAL",{**base,"locks":locks,"weeklyShifts":weekly})]
        for label,payload in variants:
            result=solve(payload,relax={"one_day_rest"})
            print("MANUAL_COMPARE",label,"EXACT",result.get("ok"),"STATUS",result.get("status"),flush=True)
            if not result.get("ok"):
                fallback=solve(payload,relax={"one_day_rest","fallback"})
                print("MANUAL_COMPARE",label,"FALLBACK",fallback.get("ok"),"STATUS",fallback.get("status"),flush=True)
                if not fallback.get("ok"):continue
                result=fallback
            sched=result["schedule"]
            hours={p:12*sum(v in ("6","18") for v in sched[p]) for p in PEOPLE}
            issues=validate({**payload,"schedule":sched},allow_shortfall=True)
            rest_issues=[x for x in issues if "tylko jeden dzień wolny" in x or "poniedziałek po sobocie" in x]
            other=[x for x in issues if x not in rest_issues]
            self.assertEqual(other,[])
            self.assertEqual(validate({**payload,"schedule":sched},allow_one_day_rest=True,allow_shortfall=True),[])
            daily=[];weekly_bad=[]
            for p in PEOPLE:
                shifts=[]
                for d,v in enumerate(sched[p]):
                    if v in ("6","18"):
                        a=datetime(2026,11,d+1,int(v));shifts.append((a,a+timedelta(hours=12)))
                for a,b in zip(shifts,shifts[1:]):
                    gap=(b[0]-a[1]).total_seconds()/3600
                    if gap<11:daily.append((p,a[0].day,b[0].day,gap))
                for d in range(1,25):
                    left=datetime(2026,11,d);right=left+timedelta(days=7)
                    intervals=[(left,left)]+[(max(left,a),min(right,b)) for a,b in shifts if a<right and b>left]+[(right,right)]
                    intervals.sort()
                    gap=max((intervals[i+1][0]-intervals[i][1]).total_seconds()/3600 for i in range(len(intervals)-1))
                    if gap<35:weekly_bad.append((p,d,round(gap,1)))
            print("MANUAL_HOURS",label,hours,flush=True)
            print("MANUAL_REST",label,"ONE_DAY",len(rest_issues),"DAILY_UNDER_11",daily,"ROLLING_WEEK_UNDER_35",len(weekly_bad),weekly_bad[:30],flush=True)
            self.assertEqual(daily,[])
    def test_validator_rejects_invalid(self):
        data={"year":2026,"month":11,"nominal":160,"schedule":{p:[""]*30 for p in PEOPLE}}
        issues=validate(data)
        self.assertTrue(issues)
        self.assertTrue(any("wymagane dokładnie" in i for i in issues))
if __name__=="__main__":
    unittest.main(verbosity=2)
